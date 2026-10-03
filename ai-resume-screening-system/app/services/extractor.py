import os
import re
import logging

logger = logging.getLogger(__name__)

def clean_extracted_text(text: str) -> str:
    """
    Cleans and normalizes extracted resume text.
    Handles irregular unicode characters, bullets, multiple newlines, and whitespaces.
    """
    if not text:
        return ""
        
    # Replace non-breaking spaces and zero-width spaces
    text = text.replace('\xa0', ' ').replace('\u200b', ' ').replace('\ufeff', ' ')
    
    # Replace custom bullet symbols and dashes
    text = re.sub(r'[\uf0b7\uf0a7\uf076\u2022\u2023\u25e6\u2043\u2219\u25cf\u25cb\u25aa\u25ab]', '\n- ', text)
    text = text.replace('\u2013', '-').replace('\u2014', '-').replace('\u2018', "'").replace('\u2019', "'")
    text = text.replace('\u201c', '"').replace('\u201d', '"')
    
    # Remove control characters except standard whitespace / newlines
    text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f]', '', text)
    
    # Normalize tabs to spaces
    text = text.replace('\t', ' ')
    
    # Clean up multi-spaces per line
    lines = [re.sub(r' +', ' ', line.strip()) for line in text.split('\n')]
    
    # Remove excessive blank lines (limit to max 2 consecutive newlines)
    cleaned_lines = []
    blank_count = 0
    for line in lines:
        if not line:
            blank_count += 1
            if blank_count <= 1:
                cleaned_lines.append('')
        else:
            blank_count = 0
            cleaned_lines.append(line)
            
    cleaned_text = '\n'.join(cleaned_lines).strip()
    cleaned_text = re.sub(r'\n{3,}', '\n\n', cleaned_text)
    return cleaned_text


def _extract_pdf_native(file_path: str) -> str:
    """Attempts native PDF text extraction using pdfminer and PyPDF2."""
    extracted_text = ""
    
    # 1. Try pdfminer.six first (higher accuracy with layout and font encodings)
    try:
        from pdfminer.high_level import extract_text as pdfminer_extract
        text = pdfminer_extract(file_path)
        if text and len(text.strip()) > 30:
            return text
    except Exception as e:
        logger.warning(f"pdfminer extraction failed for {file_path}: {e}")

    # 2. Fallback to PyPDF2
    try:
        import PyPDF2
        with open(file_path, 'rb') as f:
            reader = PyPDF2.PdfReader(f)
            pages_text = []
            for page in reader.pages:
                t = page.extract_text()
                if t:
                    pages_text.append(t)
            extracted_text = "\n".join(pages_text)
            if extracted_text and len(extracted_text.strip()) > 30:
                return extracted_text
    except Exception as e:
        logger.warning(f"PyPDF2 extraction failed for {file_path}: {e}")

    return extracted_text


def _extract_pdf_ocr(file_path: str) -> str:
    """Fallback OCR extraction for scanned PDF documents using pdf2image and pytesseract."""
    try:
        from pdf2image import convert_from_path
        import pytesseract
        
        logger.info(f"Triggering OCR fallback for scanned PDF: {file_path}")
        images = convert_from_path(file_path, first_page=1, last_page=5) # Process up to first 5 pages
        ocr_texts = []
        for i, image in enumerate(images):
            text = pytesseract.image_to_string(image)
            if text:
                ocr_texts.append(text)
        return "\n".join(ocr_texts)
    except Exception as e:
        logger.warning(f"OCR extraction failed for {file_path} (Tesseract/Poppler may not be installed): {e}")
        return ""


def _extract_docx(file_path: str) -> str:
    """Extracts text from DOCX files including paragraphs and tables."""
    try:
        import docx
        doc = docx.Document(file_path)
        full_text = []
        
        # Extract standard paragraphs
        for para in doc.paragraphs:
            if para.text.strip():
                full_text.append(para.text)
                
        # Extract tables (e.g. skills matrix, work history tables)
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    full_text.append(" | ".join(row_text))
                    
        return "\n".join(full_text)
    except Exception as e:
        logger.error(f"DOCX extraction error for {file_path}: {e}")
        raise ValueError(f"Failed to read DOCX file: {e}")


def extract_resume_text(file_path: str) -> str:
    """
    Main extraction interface.
    Extracts, cleans, and normalizes text from PDF or DOCX resume documents.
    Automatically triggers OCR fallback for image-based/scanned PDFs.
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")
        
    ext = os.path.splitext(file_path)[1].lower()
    raw_text = ""
    
    if ext == '.pdf':
        raw_text = _extract_pdf_native(file_path)
        
        # If native PDF extraction returned very little or no text (e.g. scanned image), try OCR
        words = raw_text.split()
        if len(words) < 25:
            logger.info(f"Native extraction yield is low ({len(words)} words). Attempting OCR.")
            ocr_text = _extract_pdf_ocr(file_path)
            if ocr_text and len(ocr_text.split()) > len(words):
                raw_text = ocr_text
                
    elif ext == '.docx':
        raw_text = _extract_docx(file_path)
    else:
        raise ValueError(f"Unsupported file format '{ext}'. Only .pdf and .docx are supported.")
        
    cleaned = clean_extracted_text(raw_text)
    return cleaned
